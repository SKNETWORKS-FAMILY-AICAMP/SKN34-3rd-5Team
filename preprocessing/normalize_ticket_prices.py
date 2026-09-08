"""
`구장정보.xlsx`의 `Ticket Prices` 시트(10구단 좌석 가격, 740행)를 전처리 스크립트.

배경 (2026-09-08):
- 기존 `kbo_ticket_policy_structured.csv`(244행)는 예매 오픈시각·매수 한도만 다루고
  실제 좌석 가격은 전혀 없었다. `구장정보.xlsx`의 다른 16개 시트와 마찬가지로
  Ticket Prices 시트도 지금까지 어떤 preprocessed CSV로도 옮겨진 적이 없었다.
- 좌석 등급명(zone_code/zone_name_ko)이 팀마다 이름형(프리미엄석 등)/색상형
  (BLUE/IVORY/NAVY 등)/티어형(TIER_1~4)으로 제각각이라 공통 스키마로 표준화하려면
  시간이 오래 걸린다. 표준화 없이 원본 필드를 그대로 옮기는 방식으로 결정했다
  (RAG는 team_code+zone_name_ko+day_type+customer_type 조합으로 자유텍스트 검색하면
  되므로 무리해서 표준화하지 않아도 됨).

QA 중 확인한 주의사항 (전처리는 하되 챗봇 답변 로직/문서에 아래를 반드시 반영할 것):

1. **NC 50행은 고정 시즌가격이 아니라 특정 경기 1건의 동적가격 스냅샷이다**
   (`day_type=SAMPLE_EVENT_2026-09-08`, `valid_from=valid_to=2026-09-08`,
   `discount_condition="시즌 고정가격 아님"`). NC는 경기별로 가격이 달라지는
   다이나믹 프라이싱 구조라, 다른 9개 구단처럼 WEEKDAY/WEEKEND_HOLIDAY 고정가가
   아니다. 챗봇이 NC 가격을 답할 때는 "2026-09-08 특정 경기 기준 예시가격이며
   실제 요금은 경기별로 달라질 수 있습니다" 같은 디스클레이머를 반드시 붙여야
   한다. (그대로 두면 마치 고정가처럼 안내될 위험이 있음)
2. **PARTIAL 5행은 전부 NC 스카이박스/불펜/라운드테이블석** — 표시금액은 있으나
   1인 단가인지 구역 전체 총액인지 목록만으로 불명확(`status=PARTIAL`). 그대로
   노출하되 "결제 단계에서 단위 재확인 필요"라는 원본 안내를 답변에 포함해야 함.
3. **`group_size` 컬럼이 숫자(6, 4)와 텍스트("무료 대상 증빙 필요")가 섞여 있다.**
   숫자로 파싱해서 쓰는 로직이 있다면 이 컬럼을 숫자 전용으로 가정하면 안 된다.
4. **`price_krw=0`인 4행은 오류가 아니라 진짜 무료 요금**(LG 엘린이 그린석 외야,
   키움 미취학아동/리틀야구 회원 외야 상단 — 전부 `discount_condition`/customer_type에
   무료 대상 조건이 명시돼 있음). 0원을 그대로 노출하되 무료 대상 조건 문구를
   같이 보여줘야 오해가 없다.

컬럼은 원본 그대로 유지하고 `evidence_type`(공통 어휘 정규화, 전부 OFFICIAL)만 추가한다.
"""
import openpyxl
import pandas as pd
from pathlib import Path

SCRIPT_DIR = Path(__file__).parent
RAW_DIR = SCRIPT_DIR / ".." / "data" / "raw"
OUT_DIR = SCRIPT_DIR / ".." / "data" / "preprocessed"
SRC_XLSX = RAW_DIR / "구장정보.xlsx"
OUT_CSV = OUT_DIR / "구장티켓가격.csv"

VALID_TEAM_CODES = {"LG", "DOOSAN", "KIWOOM", "SSG", "KT", "HANWHA", "SAMSUNG", "KIA", "LOTTE", "NC"}
VALID_STADIUM_CODES = {"JAMSIL", "GOCHEOK", "MUNHAK", "SUWON", "DAEJEON", "DAEGU", "GWANGJU", "SAJIK", "CHANGWON"}


def process_ticket_prices(wb):
    ws = wb["Ticket Prices"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    df = pd.DataFrame(rows)

    bad_team = set(df["team_code"]) - VALID_TEAM_CODES
    bad_stadium = set(df["stadium_code"]) - VALID_STADIUM_CODES
    if bad_team:
        print(f"[경고] 알 수 없는 team_code: {bad_team}")
    if bad_stadium:
        print(f"[경고] 알 수 없는 stadium_code: {bad_stadium}")

    df["evidence_type"] = "OFFICIAL"

    col_order = [
        "season", "team_code", "stadium_code", "zone_code", "zone_name_ko",
        "price_tier", "day_type", "customer_type", "group_size", "price_krw",
        "valid_from", "valid_to", "discount_condition", "evidence_type",
        "source_id", "verified_at", "status", "notes",
    ]
    df = df[col_order]
    df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
    print(f"저장: {OUT_CSV} ({len(df)}행)")
    print(df["team_code"].value_counts())


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    process_ticket_prices(wb)
