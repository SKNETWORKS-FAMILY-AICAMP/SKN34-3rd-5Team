"""
`구장정보.xlsx`의 `Seat Zones` 시트(10구단 좌석 구역 상세, 192행)를 전처리하는 스크립트.

배경 (2026-09-08 밤, 3차 추가):
- `구장정보.xlsx`의 17개 시트 중 Facilities(편의시설), Transport(교통·주차),
  Ticket Prices(가격), Policies/Team Policies(정책)는 이미 전처리를 끝냈는데,
  Seat Zones(구역 상세, 192행)만 "시간 관계상 저우선순위로 이월" 상태로 남아있었다.
  형준님이 나머지 미결 항목을 마저 끝내자고 해서 이번에 처리한다.
- Ticket Prices와 마찬가지로 zone_code/zone_name_ko가 팀마다 이름형(프리미엄석 등)/
  색상형(BLUE/GREEN 등)/기능형(SKYBOX/BBQ 등)으로 제각각이라 공통 스키마로
  표준화하지 않고 원본 필드를 그대로 옮긴다. RAG는 team_code+zone_name_ko(또는
  zone_code) 조합으로 자유텍스트 검색하면 된다.

QA 중 확인한 사항:
- `level`(층수)은 192행 중 170행이 비어있다(원본에 값이 없는 정상 케이스 — 옥외/전체층
  구역 등 층 구분이 의미 없는 구역으로 추정). `side`(측면)도 50행이 비어있고,
  `group_size`(그룹석 인원)는 150행이 비어있다 — 전부 원본에 애초에 없는 값이라
  결측이 아니라 "해당 없음"으로 취급해야 한다. NULL을 파싱 오류로 취급하지 말 것.
- `status=PARTIAL`은 2행뿐(삼성 대구 1행, KIA 광주 1행) — Ticket Prices만큼 불확실성이
  크지 않다. 상세 사유는 각 행의 `description` 필드에 원본 그대로 들어있으니 그대로
  노출하면 된다(Ticket Prices처럼 별도 챗봇 응답 규칙을 추가할 정도의 이슈는 아님).
- `accessible=Y`인 휠체어석/동반석 등은 원본 그대로 유지 — 별도 가공 없음.

컬럼은 원본 그대로 유지하고 `evidence_type`(공통 어휘 정규화, 전부 OFFICIAL)만 추가한다.
"""
import openpyxl
import pandas as pd
from pathlib import Path

SCRIPT_DIR = Path(__file__).parent
RAW_DIR = SCRIPT_DIR / ".." / "data" / "raw"
OUT_DIR = SCRIPT_DIR / ".." / "data" / "preprocessed"
SRC_XLSX = RAW_DIR / "구장정보.xlsx"
OUT_CSV = OUT_DIR / "구장좌석구역.csv"

VALID_TEAM_CODES = {"LG", "DOOSAN", "KIWOOM", "SSG", "KT", "HANWHA", "SAMSUNG", "KIA", "LOTTE", "NC"}
VALID_STADIUM_CODES = {"JAMSIL", "GOCHEOK", "MUNHAK", "SUWON", "DAEJEON", "DAEGU", "GWANGJU", "SAJIK", "CHANGWON"}


def process_seat_zones(wb):
    ws = wb["Seat Zones"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    df = pd.DataFrame(rows)

    bad_team = set(df["team_code"]) - VALID_TEAM_CODES
    bad_stadium = set(df["stadium_code"]) - VALID_STADIUM_CODES
    if bad_team:
        print(f"[경고] 알 수 없는 team_code: {bad_team}")
    if bad_stadium:
        print(f"[경고] 알 수 없는 stadium_code: {bad_stadium}")

    dup = df.duplicated(subset=["team_code", "stadium_code", "zone_code"]).sum()
    if dup:
        print(f"[경고] team_code+stadium_code+zone_code 중복 {dup}건 발견")

    df["evidence_type"] = "OFFICIAL"

    col_order = [
        "season", "team_code", "stadium_code", "zone_code", "zone_name_ko",
        "level", "side", "seat_type", "group_size", "accessible", "description",
        "evidence_type", "source_id", "verified_at", "status",
    ]
    df = df[col_order]
    df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
    print(f"저장: {OUT_CSV} ({len(df)}행)")
    print(df["team_code"].value_counts())
    print("\nstatus 분포:")
    print(df["status"].value_counts())


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    process_seat_zones(wb)
