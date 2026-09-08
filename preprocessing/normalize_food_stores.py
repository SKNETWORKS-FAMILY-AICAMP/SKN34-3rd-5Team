"""
공식 구장 먹거리·부가콘텐츠(`구장먹거리,컨텐츠.xlsx`) 전처리 스크립트

배경 (2026-09-08 점검 중 발견):
    이 xlsx 파일은 "구장 내부 매점"(Food Stores, 284건)과 "분리하기" 정책의 두 축 중 하나인데,
    다른 모든 데이터(카카오 external_places.csv, 편의시설, 좌석/주차/버스 등)는 data/preprocessed/에
    CSV로 전처리돼 있는 반면 이 파일만 raw xlsx 상태 그대로 남아있었다 — RAG 청킹 전에 반드시
    끝내야 할 누락 작업.
    같은 파일 안의 "Attractions" 시트(68건, 구장 내부 키즈존·포토존·체험존 등 부가 콘텐츠)도
    지금까지 어떤 문서에도 언급되지 않은 채 방치돼 있었다 — 이번에 같이 전처리한다.

발견한 이슈와 처리 방식:
    - Food Stores·Attractions 두 시트 모두 Evidence Type이 OFFICIAL_GUIDE_MAP, OFFICIAL_GUIDE_PAGE,
      OFFICIAL_TEAM_PRESS 등 세분화된 값으로 돼 있다. 프로젝트 공통 evidence_type 어휘(OFFICIAL/
      DERIVED/THIRD_PARTY_API/UNOFFICIAL)와 다른 이름 체계라, 반입물품·재입장 파일에서 했던 것과
      동일하게 `evidence_type=OFFICIAL`(공통 어휘)로 정규화하고, 원래 세분값은 `evidence_subtype`
      컬럼에 그대로 보존한다. 필터링 로직은 `evidence_type`만 보면 되고, 출처 세부가 필요하면
      `evidence_subtype`을 추가로 본다.
    - stadium_code는 원본에 이미 프로젝트 표준 9개 코드(JAMSIL/GOCHEOK/MUNHAK/SUWON/DAEJEON/
      DAEGU/GWANGJU/SAJIK/CHANGWON)로 들어있어 별도 매핑이 필요 없었다(정규화 스크립트가 필요했던
      팀원 보완자료 2종과 다른 점).
    - Record ID(F001.., A001..)는 중복 없이 고유함을 확인.
    - Status는 Food Stores 282건 CONFIRMED + 2건 PARTIAL, Attractions는 68건 전부 CONFIRMED —
      원본 값을 그대로 유지(프로젝트 공통 status 어휘와 이름이 이미 같음).
    - Coverage/Backlog 시트는 팀 내부 QA·품질 추적용 메모라 RAG에 넣지 않고 그대로 원본에 둔다.
      (Backlog가 F&B 매장별 영업시간 부족을 이미 자체적으로 추적 중 — 우리 `비공식조사_필요항목_정리.md`의
      "F&B 영업시간 우선순위 낮음" 항목과 내용이 일치, 새로운 문제는 아님)

주의:
    `구장먹거리,컨텐츠.xlsx`가 로컬에서 Excel로 열려있으면(잠금 파일 ~$구장먹거리,컨텐츠.xlsx 존재)
    읽기는 되지만 원본에 다시 쓰지는 않는다 — 이 스크립트도 원본은 절대 수정하지 않고 읽기 전용으로만
    연다. 출력은 전부 data/preprocessed/ 새 파일.

---
2026-09-08 밤 추가 — `backlog 보충.xlsx`(팀원 3차 재검증) `신규확보데이터` No.10 반영:
    DAEJEON(한화) Food Stores에 "더본코리아 입점 브랜드 구역"이라는 PARTIAL 요약 행 1개만
    있었는데(브랜드명 비공개 상태), 더본코리아 공식 보도자료에서 3루측 8개 브랜드 실명
    (빽다방빵연구소·새마을식당·역전우동·빽보이피자·한신포차·연돈볼카츠·백스비어·고투웍)이
    확인됐다. xlsx 원본에는 없는 정보라 `apply_daejeon_deoban_enrichment()` 오버레이로
    xlsx 파싱 이후 적용 — 기존 요약 행(F283)을 지우고 브랜드별 개별 행 8개로 대체한다.
    출처 자료일이 2025-03-04(2026시즌 유지 여부 재확인 필요)라 status는 CONFIRMED로 올리지
    않고 PARTIAL 유지. 개별 메뉴가 확인된 3개 브랜드(빽다방빵연구소/역전우동/백스비어)만
    구체적 메뉴를 적고 나머지 5개는 "세부 미공개"로 남긴다.

    **source_id 주의**: `구장먹거리,컨텐츠.xlsx`의 Sources 시트는 S01~S28까지만 등록돼 있고
    이번에 반영한 더본코리아 공식 보도자료(backlog 보충.xlsx 경유)는 그 표에 없는 완전히
    새로운 출처다. 처음엔 실수로 기존 S22(전혀 다른 출처 — "4월 홈경기일 포토존" 관련)를
    재사용할 뻔했다가 Sources 시트를 대조하는 과정에서 발견해 **S29로 정정**했다. S29는
    아직 원본 xlsx의 Sources 시트에는 등록되어 있지 않은 번호이므로, 원본 xlsx를 다음에
    갱신할 기회가 있으면 Sources 시트에도 S29로 공식 등록해 번호가 어긋나지 않게 할 것.

사용법:
    pip install openpyxl pandas
    python normalize_food_stores.py
"""

import openpyxl
import pandas as pd
from pathlib import Path

SCRIPT_DIR = Path(__file__).parent
RAW_DIR = SCRIPT_DIR / ".." / "data" / "raw"
OUT_DIR = SCRIPT_DIR / ".." / "data" / "preprocessed"

SRC_XLSX = RAW_DIR / "구장먹거리,컨텐츠.xlsx"

FOOD_OUT = OUT_DIR / "구장먹거리_공식매점.csv"
ATTR_OUT = OUT_DIR / "구장부가콘텐츠_공식.csv"

VALID_STADIUM_CODES = {
    "JAMSIL", "GOCHEOK", "MUNHAK", "SUWON", "DAEJEON",
    "DAEGU", "GWANGJU", "SAJIK", "CHANGWON",
}


def _sheet_to_df(wb, sheet_name, header_row):
    ws = wb[sheet_name]
    header = [c.value for c in ws[header_row]]
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, len(header) + 1)]
        if all(v is None for v in row):
            continue
        rows.append(row)
    return pd.DataFrame(rows, columns=header)


def _check_stadium_codes(df, label):
    bad = set(df["Stadium Code"].dropna().unique()) - VALID_STADIUM_CODES
    if bad:
        print(f"[WARN] {label}: 표준 stadium_code가 아닌 값 발견 — {bad} (수동 확인 필요)")


DEOBAN_BRANDS = [
    ("빽다방빵연구소", "최강한화 홈런볼빵 등 베이커리류"),
    ("새마을식당", "다양한 메뉴(브랜드 대표 메뉴, 세부 미공개)"),
    ("역전우동", "컵 우동·비빔면"),
    ("빽보이피자", "다양한 메뉴(브랜드 대표 메뉴, 세부 미공개)"),
    ("한신포차", "다양한 메뉴(브랜드 대표 메뉴, 세부 미공개)"),
    ("연돈볼카츠", "다양한 메뉴(브랜드 대표 메뉴, 세부 미공개)"),
    ("백스비어", "홈플레이트 세트메뉴 등"),
    ("고투웍", "다양한 메뉴(브랜드 대표 메뉴, 세부 미공개)"),
]


def apply_daejeon_deoban_enrichment(df):
    """더본코리아 8개 브랜드 실명 반영 (backlog 보충.xlsx 신규확보데이터 No.10).
    기존 '더본코리아 입점 브랜드 구역' 요약 행(store_facility로 식별)을 지우고
    브랜드별 개별 행으로 대체한다."""
    mask = (df["stadium_code"] == "DAEJEON") & (df["store_facility"] == "더본코리아 입점 브랜드 구역")
    removed = mask.sum()
    if removed == 0:
        print("[정보] '더본코리아 입점 브랜드 구역' 요약 행을 찾지 못함 — 원본이 이미 바뀐 것일 수 있으니 확인 필요")
        return df
    df = df[~mask].copy()

    existing_nums = df["record_id"].str.extract(r"F(\d+)")[0].dropna().astype(int)
    next_num = existing_nums.max() + 1

    new_rows = []
    for i, (name, menu) in enumerate(DEOBAN_BRANDS):
        new_rows.append({
            "record_id": f"F{next_num + i}",
            "stadium_code": "DAEJEON",
            "stadium_name": "대전 한화생명 볼파크",
            "home_team": "HANWHA",
            "store_facility": name,
            "floor": None,
            "zone_location": "3루측",
            "map_label": None,
            "menu_category_official": menu,
            "location_qty": 1,
            "operating_hours_official": None,
            "evidence_type": "OFFICIAL",
            "evidence_subtype": "OFFICIAL_PRESS_RELEASE",
            "source_id": "S29",
            "source_url": None,
            "verified_at": "2025-03-04",
            "status": "PARTIAL",
            "notes": (
                "더본코리아 공식 보도자료 기준(2025-03-04 확인, 2026시즌 유지 여부 재확인 필요). "
                "backlog 보충.xlsx 신규확보데이터 No.10 경유로 확보 — 원 보도자료 URL은 미기재."
            ),
        })
    return pd.concat([df, pd.DataFrame(new_rows)], ignore_index=True)


def process_food_stores(wb):
    df = _sheet_to_df(wb, "Food Stores", header_row=4)
    _check_stadium_codes(df, "Food Stores")

    dup = df["Record ID"].duplicated().sum()
    if dup:
        print(f"[WARN] Food Stores: Record ID 중복 {dup}건 발견")

    out = pd.DataFrame({
        "record_id": df["Record ID"],
        "stadium_code": df["Stadium Code"],
        "stadium_name": df["Stadium"],
        "home_team": df["Home Team"],
        "store_facility": df["Store / Facility"],
        "floor": df["Floor"],
        "zone_location": df["Zone / Location"],
        "map_label": df["Map Label"],
        "menu_category_official": df["Menu / Category (official)"],
        "location_qty": df["Location Qty"],
        "operating_hours_official": df["Operating Hours (official)"],
        "evidence_type": "OFFICIAL",
        "evidence_subtype": df["Evidence Type"],
        "source_id": df["Source ID"],
        "source_url": df["Source URL"],
        "verified_at": df["Verified At"],
        "status": df["Status"],
        "notes": df["Notes"],
    })
    out = apply_daejeon_deoban_enrichment(out)
    out.to_csv(FOOD_OUT, index=False, encoding="utf-8-sig")
    print(f"완료: {FOOD_OUT} ({len(out)}행)")


def process_attractions(wb):
    df = _sheet_to_df(wb, "Attractions", header_row=4)
    _check_stadium_codes(df, "Attractions")

    dup = df["Record ID"].duplicated().sum()
    if dup:
        print(f"[WARN] Attractions: Record ID 중복 {dup}건 발견")

    out = pd.DataFrame({
        "record_id": df["Record ID"],
        "stadium_code": df["Stadium Code"],
        "stadium_name": df["Stadium"],
        "home_team": df["Home Team"],
        "content_type": df["Content Type"],
        "content_name": df["Content Name"],
        "floor": df["Floor"],
        "location": df["Location"],
        "official_description": df["Official Description"],
        "operating_condition": df["Operating Condition"],
        "evidence_type": "OFFICIAL",
        "evidence_subtype": df["Evidence Type"],
        "source_id": df["Source ID"],
        "source_url": df["Source URL"],
        "verified_at": df["Verified At"],
        "status": df["Status"],
        "notes": df["Notes"],
    })
    out.to_csv(ATTR_OUT, index=False, encoding="utf-8-sig")
    print(f"완료: {ATTR_OUT} ({len(out)}행)")


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
    process_food_stores(wb)
    process_attractions(wb)
